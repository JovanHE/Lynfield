from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ── 调色板（与 HTML 文档同系）──────────────────────────────────────
INK_DEEP    = "1C2E35"   # 深墨色
INK_MID     = "3E5560"   # 中墨色
INK_SOFT    = "5C7A88"   # 柔墨色
INK_FAINT   = "8FADB8"   # 淡墨色

PARCHMENT   = "F7F3EA"   # 羊皮底色
PARCHMENT2  = "F0EBE0"   # 稍深底色

WATER_HEAD  = "2E5F7A"   # 水蓝深（表头）
WATER_DARK  = "3A7491"   # 水蓝
WATER_MID   = "5FA8C8"   # 水蓝中
WATER_PALE  = "C5DDE8"   # 水蓝浅（偶数行底）
WATER_GHOST = "E8F3F8"   # 水蓝极淡（奇数行底）

LICHEN_HEAD = "2E5C3A"   # 苔绿深（表头）
LICHEN_MID  = "5C9C6A"   # 苔绿中
LICHEN_PALE = "C5DCC9"   # 苔绿浅
LICHEN_GHOST= "EAF3EB"   # 苔绿极淡

RUST_HEAD   = "7A3A22"   # 锈橙深（表头）
RUST_MID    = "B05A35"   # 锈橙中
RUST_PALE   = "E8C9B8"   # 锈橙浅
RUST_GHOST  = "F8EDE8"   # 锈橙极淡

PURPLE_HEAD = "4A2E5C"   # 紫深（角色表头）
PURPLE_PALE = "DDD0E8"
PURPLE_GHOST= "F3EEF8"

GOLD_HEAD   = "5C4A1A"   # 金深（概览表头）
GOLD_PALE   = "EAD89A"
GOLD_GHOST  = "FAF5E4"

WHITE       = "FFFFFF"

def thin_border(color="BBBBBB"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def make_font(bold=False, color=INK_DEEP, size=10, italic=False):
    return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)

def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left_wrap():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def style_header_row(ws, row, cols, bg, fg=WHITE, bold=True, size=10):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = make_font(bold=bold, color=fg, size=size)
        cell.fill = make_fill(bg)
        cell.alignment = center()
        cell.border = thin_border("999999")

def style_data_row(ws, row, cols, bg_even, bg_odd, is_even):
    bg = bg_even if is_even else bg_odd
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        if not cell.fill or cell.fill.fgColor.rgb in ("00000000", "FFFFFFFF", WHITE):
            cell.fill = make_fill(bg)
        cell.border = thin_border("CCCCCC")
        if not cell.alignment or cell.alignment.horizontal == "general":
            cell.alignment = left_wrap()

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def set_row_height(ws, start, end, height=20):
    for r in range(start, end + 1):
        ws.row_dimensions[r].height = height

# ══════════════════════════════════════════════════════════════════
wb = Workbook()

# ── Sheet 1 : 概览 ─────────────────────────────────────────────
ws1 = wb.active
ws1.title = "概览"
ws1.sheet_view.showGridLines = False

# 大标题区
ws1.merge_cells("A1:D1")
c = ws1["A1"]
c.value = "山水人家  ·  Lynfield Homestead  ·  LYN"
c.font = Font(name="Arial", bold=True, color=GOLD_HEAD, size=16)
c.fill = make_fill(GOLD_GHOST)
c.alignment = center()
c.border = thin_border("BBBBBB")
ws1.row_dimensions[1].height = 38

ws1.merge_cells("A2:D2")
c = ws1["A2"]
c.value = "World Bible · 世界观数据库 · Draft I"
c.font = Font(name="Arial", italic=True, color=INK_FAINT, size=10)
c.fill = make_fill(GOLD_GHOST)
c.alignment = center()
ws1.row_dimensions[2].height = 18

# 总览字段
overview_data = [
    ("中文正式名", "山水人家"),
    ("英文正式名", "Lynfield Homestead"),
    ("官方简称",   "LYN（3 字母标识）"),
    ("核心理念",   "神话只做氛围与情感赋魅，不改变现实衰败的真实原因——虚实两层并行，以浪漫宿命注解现实荒芜"),
    ("神话核心意象", "水山人——无史实、无器物遗迹，纯虚构灵脉先民符号，与草木流水共生，只是自然栖居守望山水情意"),
    ("全域水系中心", "美林泉（所有主溪、支流、港湾、江心岛均在其境内）"),
    ("六大区块",   "美林泉 / 彩云天 / 香溪地 / 诗家谷 / 白沙岛 / 清水湾"),
    ("主要地标",   "奥圭智陆神庙（Augizil Temple）/ 云丘祭台（Unqiu Altar）"),
    ("核心冲突",   "现实层：工程缺陷＋失管 → 水系枯竭 → 景观荒芜\n神话层：人类破坏水系 → 灵脉中断 → 先民情意成永久遗憾"),
    ("创作原则",   "所有景观只保留：① 原生自然样貌  ② 现实衰败现状。不添加任何虚构商业规划。"),
]

for i, (k, v) in enumerate(overview_data):
    r = i + 4
    ws1.merge_cells(f"C{r}:D{r}")
    ws1[f"A{r}"].value = k
    ws1[f"A{r}"].font = make_font(bold=True, color=WHITE)
    ws1[f"A{r}"].fill = make_fill(GOLD_HEAD)
    ws1[f"A{r}"].alignment = center()
    ws1[f"A{r}"].border = thin_border()

    ws1[f"B{r}"].border = thin_border()  # spacer col

    ws1[f"C{r}"].value = v
    ws1[f"C{r}"].font = make_font(color=INK_DEEP, size=10)
    bg = GOLD_GHOST if i % 2 == 0 else PARCHMENT
    ws1[f"C{r}"].fill = make_fill(bg)
    ws1[f"C{r}"].alignment = left_wrap()
    ws1[f"C{r}"].border = thin_border()

    h = 34 if "\n" in v else 22
    ws1.row_dimensions[r].height = h

set_col_widths(ws1, [18, 2, 55, 2])

# ── Sheet 2 : 区块 ─────────────────────────────────────────────
ws2 = wb.create_sheet("区块")
ws2.sheet_view.showGridLines = False

ws2.merge_cells("A1:F1")
c = ws2["A1"]
c.value = "六大区块 — Territory Zones"
c.font = Font(name="Arial", bold=True, color=WHITE, size=13)
c.fill = make_fill(LICHEN_HEAD)
c.alignment = center()
ws2.row_dimensions[1].height = 30

headers = ["编号", "区块名", "英文名（参考）", "角色定位", "核心地标", "现状备注"]
for col, h in enumerate(headers, 1):
    ws2.cell(row=2, column=col).value = h
style_header_row(ws2, 2, len(headers), LICHEN_HEAD)
ws2.row_dimensions[2].height = 22

blocks = [
    ("01", "美林泉", "Lynspring",     "全域水系核心，灵脉所在地",    "奥圭智陆神庙 / 灵溪 / 青樾湾 / 汀芜洲", "水系枯竭，景观荒芜，神庙九柱沉寂"),
    ("02", "彩云天", "Yunsky",        "高处景观区",                  "云丘祭台",                            "祭台紫藤仍在，望霞树寄情之处"),
    ("03", "香溪地", "Fragrant Fen",  "灵芷溪输水终点，育芷灵者居所", "—（灵芷溪末端）",                    "灵芷溪断流后植被枯萎荒芜"),
    ("04", "诗家谷", "Poet's Vale",   "—（待补全）",                 "—",                                  "—"),
    ("05", "白沙岛", "White Sands",   "—（待补全）",                 "—",                                  "—"),
    ("06", "清水湾", "Clear Cove",    "—（待补全）",                 "—",                                  "—"),
]

for i, row_data in enumerate(blocks):
    r = i + 3
    for col, val in enumerate(row_data, 1):
        cell = ws2.cell(row=r, column=col)
        cell.value = val
    bg_e, bg_o = LICHEN_PALE, LICHEN_GHOST
    style_data_row(ws2, r, len(headers), bg_e, bg_o, i % 2 == 0)
    ws2.row_dimensions[r].height = 22

set_col_widths(ws2, [8, 12, 18, 28, 30, 28])

# ── Sheet 3 : 水系 ─────────────────────────────────────────────
ws3 = wb.create_sheet("水系")
ws3.sheet_view.showGridLines = False

ws3.merge_cells("A1:G1")
c = ws3["A1"]
c.value = "全域水系 — Water System"
c.font = Font(name="Arial", bold=True, color=WHITE, size=13)
c.fill = make_fill(WATER_HEAD)
c.alignment = center()
ws3.row_dimensions[1].height = 30

headers3 = ["编号", "中文名（常态）", "中文名（衰败后）", "英文名", "类型", "所属区块", "特殊意义／羁绊说明"]
for col, h in enumerate(headers3, 1):
    ws3.cell(row=2, column=col).value = h
style_header_row(ws3, 2, len(headers3), WATER_HEAD)
ws3.row_dimensions[2].height = 22

waterways = [
    ("W01", "灵溪",   "枯滩",    "Lynxi Stream / Lynsai Creek", "主溪",   "美林泉", "整片土地的灵脉核心；活水时神庙九柱共振发光；枯竭后称「枯滩」，标志灵脉中断"),
    ("W02", "灵沐溪", "—",       "Lynmu Stream",                "支流",   "美林泉", "—"),
    ("W03", "灵芷溪", "（断流）", "Lynzhi Stream",               "支流·情感纽带", "美林泉→香溪地", "唯一向外输水水道；守泉灵者与育芷灵者的情感纽带；水流不息则情意不绝；断流后两灵者永隔"),
    ("W04", "灵汐溪", "—",       "Lynxi Stream",                "支流",   "美林泉", "—"),
    ("W05", "青樾湾", "（淤积）", "Green Shade Bay",             "天然港湾", "美林泉", "热带雨林幽静风貌；水山人隐秘相会、静守心意之处"),
    ("W06", "汀芜洲", "（荒芜）", "Wild Sandbar Islet",          "江心岛", "美林泉·灵溪中游", "灵溪中游无人小岛；灵溪枯竭后野草蔓生"),
]

for i, row_data in enumerate(waterways):
    r = i + 3
    for col, val in enumerate(row_data, 1):
        cell = ws3.cell(row=r, column=col)
        cell.value = val
    style_data_row(ws3, r, len(headers3), WATER_PALE, WATER_GHOST, i % 2 == 0)
    ws3.row_dimensions[r].height = 26

set_col_widths(ws3, [8, 14, 14, 28, 16, 20, 50])

# ── Sheet 4 : 地标 ─────────────────────────────────────────────
ws4 = wb.create_sheet("地标")
ws4.sheet_view.showGridLines = False

ws4.merge_cells("A1:G1")
c = ws4["A1"]
c.value = "核心地标 — Landmarks"
c.font = Font(name="Arial", bold=True, color=WHITE, size=13)
c.fill = make_fill(RUST_HEAD)
c.alignment = center()
ws4.row_dimensions[1].height = 30

headers4 = ["编号", "中文名", "英文名", "所属区块", "现实原型描述", "神话象征", "现状"]
for col, h in enumerate(headers4, 1):
    ws4.cell(row=2, column=col).value = h
style_header_row(ws4, 2, len(headers4), RUST_HEAD)
ws4.row_dimensions[2].height = 22

landmarks = [
    ("L01", "奥圭智陆神庙",
     "Augizil Temple\nAuguizhul Shrine",
     "美林泉",
     "水池旁9根长方体立柱装置，科技感现代设计；通电时可发光，现年久失修不再发光",
     "九柱纹路为灵脉共振印记；活水同频时发光，象征灵韵与情脉相通；如今锈蚀沉寂",
     "立柱锈蚀、照明设备损毁，不再发光；周边水系已枯竭"),
    ("L02", "云丘祭台",
     "Unqiu Altar\nUncolis Shrine",
     "彩云天",
     "景观高台，植有紫藤与望霞树",
     "紫藤为「情藤」，缠绕灵韵与情愫；望霞树是两灵者遥望相守、寄愿传情之寄托",
     "紫藤仍有生机；高台年久失修有所破损"),
]

for i, row_data in enumerate(landmarks):
    r = i + 3
    for col, val in enumerate(row_data, 1):
        cell = ws4.cell(row=r, column=col)
        cell.value = val
    style_data_row(ws4, r, len(headers4), RUST_PALE, RUST_GHOST, i % 2 == 0)
    ws4.row_dimensions[r].height = 50

set_col_widths(ws4, [8, 16, 22, 12, 38, 42, 30])

# ── Sheet 5 : 人物/灵者 ────────────────────────────────────────
ws5 = wb.create_sheet("人物·灵者")
ws5.sheet_view.showGridLines = False

ws5.merge_cells("A1:F1")
c = ws5["A1"]
c.value = "人物 · 灵者 — Characters"
c.font = Font(name="Arial", bold=True, color=WHITE, size=13)
c.fill = make_fill(PURPLE_HEAD)
c.alignment = center()
ws5.row_dimensions[1].height = 30

headers5 = ["编号", "名称／称谓", "所属区块", "身份设定", "与水系的关联", "命运与宿命"]
for col, h in enumerate(headers5, 1):
    ws5.cell(row=2, column=col).value = h
style_header_row(ws5, 2, len(headers5), PURPLE_HEAD)
ws5.row_dimensions[2].height = 22

chars = [
    ("C01", "守泉灵者\n（待命名）", "美林泉", "水山人，守护美林泉水系之灵；与灵溪同频感知，泉盈则身安，泉枯则悲愁",
     "居于美林泉，与灵溪、奥圭智陆神庙九柱共鸣；灵溪是其感知土地灵韵的媒介",
     "灵溪枯竭后，守泉灵者失去感知依托；灵芷溪断流，与育芷灵者永久相隔，情意成遗憾"),
    ("C02", "育芷灵者\n（待命名）", "香溪地", "水山人，居于香溪地，以灵芷溪水滋育芷草；与芷草共生，花开为传情信物",
     "依赖灵芷溪活水；溪中芷草花开是与守泉灵者互通情意的媒介",
     "灵芷溪断流后，芷草枯萎，育芷灵者失去情意寄托；两人隔水相望，永难相见"),
    ("C03", "水山人（族群）", "全域", "无史实、无器物遗迹的纯虚构灵脉先民；与草木流水共生，感知土地灵韵",
     "水山人整体与水系灵脉共鸣；水系健康则灵韵流转，水系枯竭则灵韵消散",
     "无刻意救世使命，只是自然栖居守望山水情意；随灵脉断绝而消隐"),
]

for i, row_data in enumerate(chars):
    r = i + 3
    for col, val in enumerate(row_data, 1):
        cell = ws5.cell(row=r, column=col)
        cell.value = val
    style_data_row(ws5, r, len(headers5), PURPLE_PALE, PURPLE_GHOST, i % 2 == 0)
    ws5.row_dimensions[r].height = 55

set_col_widths(ws5, [8, 18, 12, 42, 42, 42])

# ── Sheet 6 : 羁绊与冲突 ──────────────────────────────────────
ws6 = wb.create_sheet("羁绊·冲突")
ws6.sheet_view.showGridLines = False

ws6.merge_cells("A1:E1")
c = ws6["A1"]
c.value = "羁绊 · 冲突 — Bonds & Conflicts"
c.font = Font(name="Arial", bold=True, color=WHITE, size=13)
c.fill = make_fill(INK_DEEP)
c.alignment = center()
ws6.row_dimensions[1].height = 30

headers6 = ["编号", "名称", "层次", "涉及元素", "详细描述"]
for col, h in enumerate(headers6, 1):
    ws6.cell(row=2, column=col).value = h
style_header_row(ws6, 2, len(headers6), INK_DEEP)
ws6.row_dimensions[2].height = 22

bonds = [
    ("B01", "美林泉×香溪地 · 现实羁绊", "现实层",
     "灵芷溪 / 工程缺陷 / 维护缺失",
     "灵芷溪是美林泉天然向外输水的唯一水道，也是香溪地唯一自然水源。\n小区工程缺陷＋后期维护缺失 → 灵溪干涸 → 灵芷溪断流淤积 → 香溪地失去活水，植被枯萎荒芜。"),
    ("B02", "美林泉×香溪地 · 神话羁绊", "神话层",
     "守泉灵者 / 育芷灵者 / 灵芷溪 / 芷草",
     "守泉灵者（美林泉）与育芷灵者（香溪地）互生情愫。\n灵芷溪是天然情感纽带：泉水流经溪中、滋养芷草，花开为二人传情信物，水流不息则情意不绝。\n灵溪枯竭、灵芷溪断流后，两灵者隔水相望永难相见——水断情绝，脉竭境衰。"),
    ("B03", "世界观核心冲突 · 现实层", "现实层",
     "工程缺陷 / 失管 / 水系 / 植被",
     "先天工程防水缺陷＋后期长期失管 → 美林泉水系整体枯竭 → 主溪支流陆续断流 → 整片小区景观荒芜，野生草木回潮侵占。"),
    ("B04", "世界观核心冲突 · 神话层", "神话层",
     "人类 / 水系秩序 / 灵脉 / 先民情意",
     "人类破坏自然水系秩序 → 灵脉流转中断 → 土地失韵、草木失机 → 一段先民情意沦为永久遗憾，以浪漫宿命注解现实荒芜。"),
    ("B05", "虚实并行原则", "创作准则",
     "神话叙事 / 现实叙事",
     "神话只做氛围与情感赋魅，不改变现实衰败的真实原因。虚实两层并行，互不覆盖，相互映照。\n景观描写只保留：① 原生自然样貌  ② 现实衰败现状。不添加任何虚构商业规划。"),
]

fills_e = ["E8F0F5", "F5EDE8", PARCHMENT2, "EDE8F5", GOLD_GHOST]
fills_o = ["F5FAFC", "FAF3EE", PARCHMENT, "F5F0FA", "FAF8EE"]

for i, row_data in enumerate(bonds):
    r = i + 3
    for col, val in enumerate(row_data, 1):
        cell = ws6.cell(row=r, column=col)
        cell.value = val
        cell.fill = make_fill(fills_e[i] if i % 2 == 0 else fills_o[i])
        cell.border = thin_border("CCCCCC")
        cell.font = make_font(color=INK_DEEP)
        cell.alignment = left_wrap()
    ws6.row_dimensions[r].height = 60

set_col_widths(ws6, [8, 26, 12, 32, 62])

# ── 保存 ──────────────────────────────────────────────────────
out = r"c:\Users\ASUS\WorkBuddy\20260503123218\lynfield-worldview.xlsx"
wb.save(out)
print("saved:", out)
